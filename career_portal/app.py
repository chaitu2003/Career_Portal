import base64
from tkinter import Image
from flask import Flask,jsonify,render_template,Response, request, redirect, url_for,session, flash,send_file,send_from_directory
from flask import Flask, send_file, make_response, request
from flask_pymongo import PyMongo
from werkzeug.utils import secure_filename
from json import JSONEncoder
from datetime import datetime,timedelta
import secrets
from bson import ObjectId
import io
from io import BytesIO
from datetime import datetime
from docx import Document 
import pandas as pd
import json
from PIL import Image
from pymongo import MongoClient
from bs4 import BeautifulSoup
import requests
import os
import socket
import signal
import sys
import threading
from gridfs import GridFS
import io
import tempfile
from docx.shared import Inches
import convertapi
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
app = Flask(__name__)
app.secret_key = secrets.token_hex(16)
app.json_encoder = JSONEncoder
app.secret_key = 'supersecretkey'  # Ensure you have a secret key
convertapi.api_secret = 'XFiCeo08N1yjopSC'
# Configure MongoDB
app.config['MONGO_URI'] = 'mongodb://localhost:27017/users'
client = MongoClient('mongodb://localhost:27017/')
mongo1 = client['users']  
mongo = PyMongo(app)
db = client['users']
users_collection = db['users']
fs = GridFS(db)
MANAGER_CREDENTIALS = {
    'manager1': '1234',
    'manager2': '1234',
    'manager3': '1234',
    'manager4': '1234',
    'manager5': '1234',
    'manager6': '1234',
    'manager7': '1234',
    'manager8': '1234'
}


def get_registration_counter(registration_type):
    counter_doc = mongo.db.counters.find_one_and_update(
        {'_id': f'registration_counter_{registration_type.lower()}'},
        {'$inc': {'value': 1}},
        upsert=True,
        return_document=True
    )
    return counter_doc['value']

def generate_registration_id(registration_type):
    registration_counter = get_registration_counter(registration_type)
    prefix = "CAND"
    if registration_type.lower() == "on-campus":
        registration_type_code = "OC"
    elif registration_type.lower() == "off-campus":
        registration_type_code = "OF"
    elif registration_type.lower() == "reference":
        registration_type_code = "RF"
    elif registration_type.lower() == "walk-in":
        registration_type_code = "WI"
    
    registration_id = f"{prefix}{registration_type_code}{registration_counter:03d}"
    return registration_id


# Define the route for the homepage
@app.route('/')
def index():
    return render_template('index.html')
    

# Define the route for registration
@app.route('/registration')
def registration():

    return render_template('registration.html')


# Define the route for handling registration form submission
@app.route('/register', methods=['POST'])
def register():
    

    session.clear()   
    pname = request.form.get('pname')
    
    gender=request.form.get('gender')
    dob = request.form.get('dob')
    age=request.form.get('age')
    marital_status=request.form.get('marital_status')
    phoneNumber=request.form.get('phoneNumber')
    fatherName = request.form.get('fatherName')
    fatheroccupation=request.form.get('fatheroccupation')
    fatherNumber = request.form.get('fatherNumber')
    motherName = request.form.get('motherName')
    #motherNumber = request.form.get('motherNumber')
    Permanent_address = request.form.get('Permanent_address')
    Residencial_address = request.form.get('Residencial_address')
    schoolInstituteName =request.form.get('schoolInstituteName')
    SchoolUniversity = request.form.get('SchoolUniversity')
    SchoolDegreeName = request.form.get('SchoolDegreeName')
    Percentage_s = request.form.get('Percentage_s')
    yearofpassings=request.form.get('yearofpassings')
    IntermediateInstituteName = request.form.get('IntermediateInstituteName')
    IntermediateUniversity=request.form.get('IntermediateUniversity')
    IntermediateDegreeName = request.form.get('IntermediateDegreeName')
    Percentage_i = request.form.get('Percentage_i')
    yearofpassingi=request.form.get('yearofpassingi')
    undergraduateInstituteName = request.form.get('undergraduateInstituteName')
    undergraduateUniversity= request.form.get('undergraduateUniversity')
    undergraduateDegreeName= request.form.get('undergraduateDegreeName')
    Percentage_ug = request.form.get('Percentage_ug')
    yearofpassingug=request.form.get('yearofpassingug')

    postgraduateInstituteName = request.form.get('postgraduateInstituteName')
    postgraduateUniversity= request.form.get('postgraduateUniversity')
    postgraduateDegreeName= request.form.get('postgraduateDegreeName')
    Percentage_pg = request.form.get('Percentage_pg')
    yearofpassingpg=request.form.get('yearofpassingpg')

    email = request.form.get('email')
    registration_repsonse_time=datetime.now()
    registration_date_only = registration_repsonse_time.date()
    registration_date_only_str = registration_date_only.isoformat()
    existing_user = mongo.db.users.find_one({'email': email, 'dob': dob,'pname':pname,'phoneNumber': phoneNumber})

    if existing_user:
        # If the same set of data already exists, flash an error message
        flash('Error: User with the same data already exists!', 'error')
        return redirect(url_for('registration'))

    
    registration_type = request.form.get('registration_type')
    registration_number = generate_registration_id(registration_type)
    pdf_file = request.files['pdf_file']
    profile_image1 = request.files['img']
    # profile_image2 = request.files['img1']
    mongo.save_file(profile_image1.filename, profile_image1)
    # mongo.save_file(profile_image2.filename, profile_image2)


    work_experience=request.form.get('work_experience')
    # txtBox_work_experience=request.form.get('txtBox_work_experience')
    # txtBox_work_experience_designation=request.form.get('txtBox_work_experience_designation')
    # txtBox_work_experience_timeperiod=request.form.get('txtBox_work_experience_timeperiod')

    if work_experience == 'yes':
        organization_name = request.form.get('organization_name')
        designation = request.form.get('designation')
        time_period_from = request.form.get('time_period_from')
        time_period_to = request.form.get('time_period_to')
    else:
        organization_name = None
        
        time_period_from = None
        time_period_to = None
        designation=None

    # siblings=request.form.get('siblings')
    # txtBox_siblings=request.form.get('txtBox_siblings')

    # reference=request.form.get('reference')
    # txtBox_reference=request.form.get('txtBox_reference')


    # Retrieve sibling information
    siblings = request.form.get('siblings')
    if siblings == 'yes':
        sibling_type = request.form.get('sibling_type')
        sibling_age = request.form.get('sibling_age')
        sibling_name = request.form.get('sibling_name')
        sibling_course = request.form.get('sibling_course')
        sibling_institute = request.form.get('sibling_institute')
        sibling_location = request.form.get('sibling_location')
    else:
        sibling_type = None
        sibling_age = None
        sibling_name = None
        sibling_course = None
        sibling_institute = None
        sibling_location = None
    
    registration_type = request.form.get('registration_type')
    reference_type = None
    current_employee_id=None
    current_employee_name=None
    relation_with_reference=None
    ex_employee_name=None
    ex_employee_id=None
    walkin_type=None
    college_type=None
    location=None
    
    if registration_type == 'Off-campus':
        college_type = request.form.get('offcampus_college_name')
        location = request.form.get('offcampus_location')
    elif registration_type == 'On-campus':
        college_type = request.form.get('oncampus_college_name')
        location = request.form.get('oncampus_location')
    elif registration_type == 'Reference':
        reference_type = request.form.get('reference_type')
        if reference_type =='Current Employee':
            current_employee_id=request.form.get('current_employee_id')
            current_employee_name=request.form.get('current_employee_name')
            relation_with_reference=request.form.get('relation_with_reference')
        elif reference_type =='Ex Employee':
            ex_employee_name=request.form.get('ex_employee_name')
            ex_employee_id=request.form.get('ex_employee_id')
    elif registration_type == 'Walk-In':
        walkin_type = request.form.get('walkin_type')

    user_data = {
        'pname': pname,
        # 'MiddleName' : MiddleName,
        # 'lastName' : lastName,
        'gender': gender,
        'dob': dob,
        'age': age,
        'marital_status': marital_status,
        'phoneNumber' : phoneNumber,
        'fatherName':fatherName,
        'fatheroccupation': fatheroccupation,
        'fatherNumber' : fatherNumber,
        'motherName' : motherName,
        # 'motherNumber' : motherNumber,
        'Permanent_address' : Permanent_address,
        'Residencial_address' : Residencial_address,
        'schoolInstituteName' : schoolInstituteName,
        'SchoolUniversity' : SchoolUniversity,
        'SchoolDegreeName' : SchoolDegreeName,
        'Percentage_s' : Percentage_s,
        'yearofpassings' : yearofpassings,
        'IntermediateInstituteName' : IntermediateInstituteName,
        'IntermediateUniversity' : IntermediateUniversity,
        'IntermediateDegreeName' : IntermediateDegreeName,
        'Percentage_i' : Percentage_i,
        'yearofpassingi' : yearofpassingi,
        'undergraduateInstituteName' : undergraduateInstituteName,
        'undergraduateUniversity' : undergraduateUniversity,
        'undergraduateDegreeName' : undergraduateDegreeName,
        'Percentage_ug' : Percentage_ug,
        'yearofpassingug' : yearofpassingug,
        
        'postgraduateInstituteName' : postgraduateInstituteName,
        'postgraduateUniversity' : postgraduateUniversity,
        'postgraduateDegreeName' : postgraduateDegreeName,
        'Percentage_pg' : Percentage_pg,
        'yearofpassingpg' : yearofpassingpg,

        #'username': username,
        'email': email,
        
        'registration_type': request.form.get('registration_type'),
        'registration_number': registration_number,
        'profile_image1': profile_image1.filename,  
        # 'profile_image2': profile_image2.filename,
        'pdf_file': pdf_file.read(),
        'selected': 'not_declared',

        'work_experience': request.form.get('work_experience'),
        
        'organization_name': organization_name,
        'designation': designation,
        'time_period_from': time_period_from,
        'time_period_to': time_period_to,

         'siblings': request.form.get('siblings'),
        
        'sibling_type': sibling_type,
        'sibling_age': sibling_age,
        'sibling_name': sibling_name,
        'sibling_course': sibling_course,
        'sibling_institute': sibling_institute,
        'sibling_location': sibling_location,

        #  'campus': request.form.get('campus'),

        'college_type': college_type,
        'location': location,
        'reference_type':reference_type,
        'current_employee_id':current_employee_id,
        'current_employee_name':current_employee_name,
        'relation_with_reference':relation_with_reference,
        'ex_employee_name':ex_employee_name,
        'ex_employee_id':ex_employee_id,
        'walkin_type':walkin_type,
        'registration_repsonse_time': registration_repsonse_time,
        'registration_date_only_str':registration_date_only_str,
        'register_update' :'registered',
    }
    # Save user_data to MongoDB
    result = mongo.db.users.insert_one(user_data)
    user_id = result.inserted_id

    # Redirect to user_info route with the new user's ID
    return redirect(url_for('user_info', user_id=str(user_id)))
    # return redirect(url_for('index'))

# Define the route for login
@app.route('/login')
def login():
    return render_template('login.html')

@app.route('/login', methods=['GET', 'POST'])
def login_post():
    if request.method == 'POST':
        email = request.form.get('email')
        dob = request.form.get('dob')

        # Assuming a simple check, you should use secure authentication methods
        user = mongo.db.users.find_one({'email': email, 'dob': dob})

        if user:
            # Store registration number in session
            session['registration_number'] = user['registration_number']
            return redirect(url_for('user_new'))
        else:
            return redirect(url_for('login'))
    return render_template('login.html')

@app.route('/user_new')
def user_new():
    # Assuming you want to fetch the PDF related to the logged-in user
    registration_number = session.get('registration_number')
    if registration_number:
        user = mongo.db.users.find_one({'registration_number': registration_number})
        if user and 'pdf_id' in user:
            pdf_id = user['pdf_id']
        else:
            pdf_id = None
        return render_template('user_new.html', pdf_id=pdf_id)
    else:
        return redirect(url_for('login'))


@app.route('/user_info')
def user_info():
    # Retrieve registration number from session
    registration_number = session.get('registration_number')
    if registration_number:
        user = mongo.db.users.find_one({'registration_number': registration_number})
        return render_template('user_info.html', user=user)
    else:
        return redirect(url_for('login'))


@app.route('/get_file/<filename>')
def get_file(filename):
    file_data = mongo.db.fs.files.find_one({'filename': filename})
    if file_data:
        file_id = file_data['_id']  # Get the ObjectId of the file
        return send_file(
            io.BytesIO(mongo.db.fs.chunks.find_one({'files_id': file_id})['data']),
            mimetype='image/jpeg'
        )
    else:
        return 'File not found'
    
#define route for displaying image from the mongodb which is uploded by the user during registration
@app.route('/image/<filename>')
def get_image(filename):
    image = fs.find_one({'filename': filename})
    if not image:
        return 'Image not found!', 404
    return send_file(io.BytesIO(image.read()), mimetype='image/jpeg')

# Define the route for the admin login page
@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        if username == 'admin' and password == '123':
            session['admin_logged_in'] = True
            flash('Login successful!', 'success')
            return redirect(url_for('admin'))
        else:
            flash('Invalid credentials. Please try again.', 'error')
    return render_template('admin_login.html')

def admin_login_required(func):
    def decorator(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin_login'))
        return func(*args, **kwargs)
    decorator.__name__ = func.__name__  # Set the name of the decorator function
    return decorator

# Define the route for the admin page
@app.route('/admin')
@admin_login_required
def admin():
    return render_template('admin.html')

# Define the route for selecting users
@app.route('/select_user', methods=['POST'])
@admin_login_required
def select_user():
    registration_number = request.form.get('registration_number')
    selected = request.form.get('selected')
    mongo.db.users.update_one({'registration_number': registration_number}, {'$set': {'selected': selected}})
    flash('User selection updated successfully!', 'success')  # Flash message
    return redirect(url_for('admin'))

@app.route('/download_pdf/<registration_number>', methods=['GET'])
# @admin_login_required
# def download_pdf(registration_number):
#     user = mongo.db.users.find_one({'registration_number': registration_number})
#     if user and 'pdf_file' in user:
#         pdf_data = user['pdf_file']
#         filename = f"{user['registration_number']}.pdf"
#         response = Response(pdf_data, content_type='application/pdf')
#         response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
#         return response
#     else:
#         flash('PDF not found for this user!', 'error')
#         return redirect(url_for('admin'))

# def download_pdf(registration_number):
#     try:
#         # Construct the path to the PDF file
#         file_path = f"/path/to/pdfs/{registration_number}_pdf.docx"
        
#         # Ensure the file exists
#         if not os.path.exists(file_path):
#             return "File not found", 404
        
#         # Serve the file with inline Content-Disposition
#         return send_file(file_path, as_attachment=False, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', attachment_filename=f"{registration_number}_pdf.docx")
    
#     except Exception as e:
#         return str(e), 500

def download_pdf(registration_number):
    try:
        # Find the user document
        user = db.users.find_one({'registration_number': registration_number})
        if not user or 'pdf_file' not in user:
            return "File not found", 404

        # Retrieve the binary PDF data from the user document
        pdf_data = user['pdf_file']

        # Create a response object and set headers for inline display
        response = make_response(pdf_data)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'inline; filename={registration_number}_resume.pdf'

        return response

    except Exception as e:
        return str(e), 500



# @app.route('/download_pdf/<registration_number>', methods=['GET'])
# @admin_login_required
# def download_pdf(registration_number):
#     user = mongo.db.users.find_one({'registration_number': registration_number})
#     if user and 'pdf_file' in user:
#         pdf_data = user['pdf_file']
#         filename = f"{user['registration_number']}.pdf"
#         response = Response(pdf_data, content_type='application/pdf')
#         response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
#         return response
#     else:
#         flash('PDF not found for this user!', 'error')
#         return redirect(url_for('admin'))
# Define the route for displaying selected user details in tabular form

@app.route('/selected_user_details')
@admin_login_required
def selected_user_details():
    # Fetch selected users from the database
    selected_users = mongo.db.users.find({'selected': 'selected'})
    return render_template('selected_user_details.html', selected_users=selected_users)

@app.route('/select_user_form')
@admin_login_required
def select_user_form():
    return render_template('select_user_form.html')

@app.route('/user_details_form')
@admin_login_required
def user_details_form():
    return render_template('user_details_form.html')   

def director_login_required(func):
    def decorator(*args, **kwargs):
        if not session.get('director_logged_in'):
            return redirect(url_for('director_login'))
        return func(*args, **kwargs)
    decorator.__name__ = func.__name__  # Set the name of the decorator function
    return decorator

def manager_login_required(func):
    def decorator(*args, **kwargs):
        if not session.get('manager_logged_in'):
            return redirect(url_for('manager_login'))
        return func(*args, **kwargs)
    decorator.__name__ = func.__name__  # Set the name of the decorator function
    return decorator

@app.route('/director/login', methods=['GET', 'POST'])
def director_login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        if username == 'admin3' and password == '12345':
            session['director_logged_in'] = True
            flash('Login successful!', 'success')
            return redirect(url_for('director'))
        else:
            flash('Invalid credentials. Please try again.', 'error')
    return render_template('director_login.html')

@app.route('/manager/login', methods=['GET', 'POST'])
def manager_login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        if username in MANAGER_CREDENTIALS and password == MANAGER_CREDENTIALS[username]:
            session['manager_logged_in'] = True
            session['manager_username'] = username  # Store the manager's username in the session
            flash('Login successful!', 'success')
            return redirect(url_for('manager'))
        else:
            flash('Invalid credentials. Please try again.', 'error')
    return render_template('manager_login.html')

@app.route('/director', methods=['GET', 'POST'])
@director_login_required
def director():
    if request.method == 'POST':
        # Handle form submission and store data in the database
        registration_number = request.form.get('registration_number')
        # Get user details from the database using registration_number
        user = mongo.db.users.find_one({'registration_number': registration_number})
        if user:
            # Render a template with user details
            return render_template('dir_user_details2.html', user=user)
        else:
            flash('User not found!', 'error')
            return redirect(url_for('director'))
    return render_template('director.html')

@app.route('/manager', methods=['GET', 'POST'])
@manager_login_required
def manager():
    if request.method == 'POST':
        # Handle form submission and store data in the database
        registration_number = request.form.get('registration_number')
        # Get user details from the database using registration_number
        user = mongo.db.users.find_one({'registration_number': registration_number})
        if user:
            # Render a template with user details
            return render_template('user_details2.html', user=user)
        else:
            flash('User not found!', 'error')
            return redirect(url_for('manager'))
    return render_template('manager.html')


@app.route('/responses')
# @admin_login_required
def responses():
    return render_template('responses.html')

# Define a route to handle form submission and display responses
@app.route('/show_responses', methods=['POST'])
@admin_login_required
def show_responses():
    registration_number = request.form.get('registration_number')
    user = mongo.db.users.find_one({'registration_number': registration_number})
    
    if user:
        # Render a template with user details and responses
        return render_template('responses_display.html', user=user)
    else:
        flash('User not found!', 'error')
        return redirect(url_for('responses'))
@app.route('/store_director_details', methods=['POST'])
def store_director_details():
    registration_number = request.form.get('registration_number')
    communication = request.form.get('communication')
    Technical_Knowledge = request.form.get('Technical_Knowledge')
    Practical_Knowledge = request.form.get('Practical_Knowledge')
    Overall_Assessment = request.form.get('Overall_Assessment')
    comment = request.form.get('comment')
    Selection=request.form.get('Selection')
    if Selection == 'selected':
        selected_status = 'selected'
    else:
        selected_status = 'Not Selected'
    director_repsonse_time=datetime.now()
    director_date_only = director_repsonse_time.date()
    director_date_only_str = director_date_only.isoformat()
    mongo.db.users.update_one({'registration_number': registration_number}, 
                              {'$set': {'director_dropdown1': communication,
                                        'director_dropdown2': Technical_Knowledge,
                                        'director_dropdown3': Practical_Knowledge,
                                        'director_dropdown4': Overall_Assessment,
                                        'director_comment': comment,
                                        'director_selection': Selection,
                                        'director_repsonse_time': director_repsonse_time,
                                        'director_date_only_str':director_date_only_str,
                                        'selected': selected_status,
                                        }})

    flash('User details stored successfully!', 'success')  # Flash message
    return redirect(url_for('director'))     

@app.route('/store_manager_details', methods=['POST'])
def store_manager_details():
    registration_number = request.form.get('registration_number')
    communication = request.form.get('communication')
    Technical_Knowledge = request.form.get('Technical_Knowledge')
    Practical_Knowledge = request.form.get('Practical_Knowledge')
    Overall_Assessment = request.form.get('Overall_Assessment')
    comment = request.form.get('comment')
    Selection=request.form.get('Selection')
    manager_repsonse_time=datetime.now()
    manager_username = session.get('manager_username')
    # Update user details in the database
    mongo.db.users.update_one({'registration_number': registration_number}, 
                              {'$set': {'manager_dropdown1': communication,
                                        'manager_dropdown2': Technical_Knowledge,
                                        'manager_dropdown3': Practical_Knowledge,
                                        'manager_dropdown4': Overall_Assessment,
                                        'manager_comment': comment,
                                        'manager_selection': Selection,
                                        'manager_repsonse_time':manager_repsonse_time,
                                        'manager_username':manager_username,
                                        
                                        }})

    flash('User details stored successfully!', 'success')  # Flash message
    return redirect(url_for('manager')) 
    
@app.route('/director_selection', methods=['GET'])
def director_selection():
    return render_template('director_selection.html')  
     

@app.route('/display_director_selection', methods=['POST'])
def display_director_selection():
    selection = request.form.get('selection')
    users = mongo.db.users.find({'director_selection': selection}).sort('registration_number', 1)
    users_list = list(users)
    return render_template('director_selection_results.html', users=users_list)

@app.route('/generate_document', methods=['GET', 'POST'])
def generate_document():
    if request.method == 'POST':
        registration_number = request.form.get('registration_number')
        # Find the user by registration number
        user = mongo.db.users.find_one({'registration_number': registration_number})
        if not user:
            flash('User not found!', 'error')
            return redirect(url_for('generate_document'))

        # Load the template document from the documents directory
        doc = Document('documents/template.docx')

        # Replace placeholders in the document with actual user data
        for paragraph in doc.paragraphs:
            if '{name}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{name}', user.get('firstName', '') + ' ' + user.get('lastName', ''))
            elif '{registration_repsonse_time}' in paragraph.text:
                paragraph.text=paragraph.text.replace('{registration_repsonse_time}',user.get('registration_repsonse_time'))
            elif '{registration_number}' in paragraph.text:
                paragraph.text=paragraph.text.replace('{registration_number}',user.get('registration_number'))
            elif '{profile_image1}' in paragraph.text:
                paragraph.text=paragraph.text.replace('{profile_image1 }',user.get('profile_image1 '))  
            elif '{director_selection}' in paragraph.text:
                paragraph.text=paragraph.text.replace('{director_selection }',user.get('director_selection '))  
        # Save the modified document to a BytesIO object
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)

        # Send the modified document as a download
        return send_file(output, as_attachment=True, download_name=f"{registration_number}_pdf.docx", mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    return render_template('generate_document.html')

@app.route('/download_user_document/<user_id>', methods=['POST'])
def download_user_document(user_id):
    data = mongo.db.users.find_one({"_id": ObjectId(user_id)})
    registration_number=data.get('registration_number')

    # Check if data is found
    if not data:
        return "No data found for this registration number", 404

    # Generate PDF
    pdf_file = f"{registration_number}.pdf"
    c = canvas.Canvas(pdf_file, pagesize=letter)

    c.setFont("Helvetica", 10)
    x_offset = 100
    y_offset = 750
    line_height = 15
    if 'profile_image1' in data and data['profile_image1']:
        try:
            # Debugging: Print a portion of the base64 string
            base64_str = data['profile_image1']
            print(f"Base64 string (first 100 characters): {base64_str[:100]}")

            # Ensure the base64 string has the proper padding
            missing_padding = len(base64_str) % 4
            if missing_padding:
                base64_str += '=' * (4 - missing_padding)

            # Decode base64 string
            image_data = base64.b64decode(base64_str)

            # Debugging: Check the size of the decoded data
            print(f"Decoded image data size: {len(image_data)} bytes")

            # Create an in-memory binary stream to hold the image
            image_stream = io.BytesIO(image_data)

            # Load the image from the binary stream
            with Image.open(image_stream) as image:
                image.verify()  # Verify that it is, in fact, an image
                image = Image.open(image_stream)  # Re-open the image for drawing

                # Save the image to a temporary file to use with drawImage
                temp_image_path = 'temp_profile_image1.jpg'
                image.save(temp_image_path)

                # Draw the image on the PDF
                c.drawImage(temp_image_path, x_offset, y_offset - 100, width=62, height=62)
                y_offset -= 120  # Adjust y_offset to account for image height
        except (base64.binascii.Error, IOError, Image.UnidentifiedImageError) as e:
            print(f"Error handling image data: {e}")
            # Handle the error appropriately, such as logging or providing a default image
    c.drawString(x_offset, y_offset, f"Registration Date: {data.get('registration_repsonse_time', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Registration Number: {data.get('registration_number', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Name of Candidate: {data.get('pname', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Present Address: {data.get('Residencial_address', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Permanent Address: {data.get('Permanent_address', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Contact: {data.get('phoneNumber', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"E-mail: {data.get('email', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Marital Status: {data.get('marital_status', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Date of Birth: {data.get('dob', 'N/A')} Age: {data.get('age', 'N/A')} Gender: {data.get('gender', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Name of Father: {data.get('fatherName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Occupation of Father: {data.get('fatheroccupation', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Father Contact: {data.get('fatherNumber', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Name of Mother: {data.get('motherName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Siblings Details: {data.get('sibling_name', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Occupation: {data.get('sibling_course', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Academic Details:")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"S.S.C(10th) - Name of School/University: {data.get('schoolInstituteName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Stream: {data.get('SchoolUniversity', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Board: {data.get('SchoolDegreeName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Year of Passing: {data.get('yearofpassings', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Marks(%): {data.get('Percentage_s', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"H.S.C(12th) - Name of School/University: {data.get('IntermediateInstituteName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Stream: {data.get('IntermediateUniversity', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Board: {data.get('IntermediateDegreeName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Year of Passing: {data.get('yearofpassingi', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Marks(%): {data.get('Percentage_i', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"UG/Degree/Diploma - Name of Institute/University: {data.get('undergraduateInstituteName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"University: {data.get('undergraduateUniversity', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Degree: {data.get('undergraduateDegreeName', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Year of Passing: {data.get('yearofpassingug', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Marks(%): {data.get('Percentage_ug', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Work Experience - Company: {data.get('organization_name', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Role: {data.get('designation', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Period: {data.get('time_period_from', 'N/A')} to {data.get('time_period_to', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Aadhar No: {data.get('aadhar_no', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"PAN No: {data.get('pan_no', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Reference Check (Current) - Employee ID: {data.get('current_employee_id', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Employee Name: {data.get('current_employee_name', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Relation with Employee: {data.get('relation_with_reference', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Reference Check (Ex-Employee) - Employee Name: {data.get('ex_employee_name', 'N/A')}")
    y_offset -= line_height
    c.drawString(x_offset, y_offset, f"Selection: {data.get('director_selection', 'N/A')}")
    
    c.save()

    # Send generated PDF as a downloadable file
    # return send_file(pdf_file, as_attachment=True,mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.pdf')
    return send_file(pdf_file, as_attachment=True, download_name=f"{registration_number}.pdf", mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.pdf')
   
def convert_image_to_supported_format(image_data, format='PNG'):
    """Converts an image to a supported format and returns a BytesIO object."""
    image = Image.open(io.BytesIO(image_data))
    converted_image = BytesIO()
    image.save(converted_image, format=format)
    converted_image.seek(0)
    return converted_image
@app.route('/registered_details')
def registered_details():
    session.clear()
    users = list(mongo.db.users.find({"registration_number": {"$exists": True}}))
    for user in users:
        user['_id'] = str(user['_id'])
        user['registration_number'] = user.get('registration_number', '')
        user['pname'] = user.get('pname', '')
        user['phoneNumber'] = user.get('phoneNumber', '')
        user['email'] = user.get('email', '')
    return render_template('registered_details.html', users=users)

@app.route('/results')
def results():
    users = list(mongo.db.users.find({"director_selection": {"$exists": True}}))
    # users = list(mongo.db.users.find({"director_selection": "selected"}))
    for user in users:
        user['_id'] = str(user['_id'])
    return render_template('director_selection_results.html', users=users)

@app.route('/director_selection_redirect_selected', methods=['POST'])
def director_selection_redirect_selected():
    session['director_selection'] = 'selected'
    return redirect(url_for('director_selection_results_with_default'))

@app.route('/director_selection_redirect_not_selected', methods=['POST'])
def director_selection_redirect_not_selected():
    session['director_selection'] = 'not_selected'
    return redirect(url_for('director_selection_results_with_default'))

@app.route('/director_selection_results_with_default', methods=['GET'])
def director_selection_results_with_default():
    selection = session.get('director_selection', 'selected')
    users = list(users_collection.find({"director_selection": selection}))

    for user in users:
        user['_id'] = str(user['_id'])
        user['registration_number'] = user.get('registration_number', '')
        user['pname'] = user.get('pname', '')
        user['phoneNumber'] = user.get('phoneNumber', '')
        user['email'] = user.get('email', '')

    return render_template('director_selection_results.html', users=users, selection=selection)
    
@app.route('/download_excel', methods=['POST'])
def download_excel():
    try:
        session.clear()
        # The URL of the page to scrape
        url = request.host_url + 'director_selection_results'
        
        # Send an HTTP GET request to the URL
        response = requests.get(url)
        
        # Parse the HTML code using BeautifulSoup
        soup = BeautifulSoup(response.content, 'html.parser')
        
        # Extract the relevant information from the HTML code
        rows = []
        for row in soup.select('tbody tr'):
            cells = row.find_all('td')
            if cells:
                registration_number = cells[0].get_text(strip=True)
                name = cells[1].get_text(strip=True)
                phone = cells[2].get_text(strip=True)
                email = cells[3].get_text(strip=True)
                rows.append([registration_number, name, phone, email])
        
        # Store the information in a pandas dataframe
        df = pd.DataFrame(rows, columns=['Registration Number', 'Name', 'Phone', 'Email'])
        
        # Create a BytesIO object to write the CSV file
        output = BytesIO()
        
        # Use the DataFrame to create a CSV file
        df.to_csv(output, index=False)
        
        # Seek to the beginning of the stream
        output.seek(0)
        
        # Send the CSV file as a download
        return send_file(output, as_attachment=True, download_name='users.csv', mimetype='text/csv')
    
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/download_excel1', methods=['POST'])
def download_excel1():
    session.clear()
    try:
        
        # The URL of the page to scrape
        url = request.host_url + 'registered_details'
        
        # Send an HTTP GET request to the URL
        response = requests.get(url)
        
        # Parse the HTML code using BeautifulSoup
        soup = BeautifulSoup(response.content, 'html.parser')
        
        # Extract the relevant information from the HTML code
        rows = []
        for row in soup.select('tbody tr'):
            cells = row.find_all('td')
            if cells:
                time_stamp=cells[0].get_text(strip=True)
                registration_number = cells[1].get_text(strip=True)
                name = cells[2].get_text(strip=True)
                phone = cells[3].get_text(strip=True)
                email = cells[4].get_text(strip=True)
                date = cells[5].get_text(strip=True)
                fatherName = cells[6].get_text(strip=True)
                fatherNumber=cells[7].get_text(strip=True)
                fatheroccupation=cells[8].get_text(strip=True)
                motherName=cells[9].get_text(strip=True)
                Residencial_address=cells[10].get_text(strip=True)
                Permanent_address=cells[11].get_text(strip=True)
                schoolInstituteName=cells[12].get_text(strip=True)
                SchoolUniversity=cells[13].get_text(strip=True)
                SchoolDegreeName=cells[14].get_text(strip=True)
                Percentage_s=cells[15].get_text(strip=True)
                yearofpassings=cells[16].get_text(strip=True)
                IntermediateInstituteName=cells[17].get_text(strip=True)
                IntermediateUniversity=cells[18].get_text(strip=True)
                IntermediateDegreeName=cells[19].get_text(strip=True)
                Percentage_i=cells[20].get_text(strip=True)
                yearofpassingi=cells[21].get_text(strip=True)

                undergraduateInstituteName=cells[22].get_text(strip=True)
                undergraduateUniversity=cells[23].get_text(strip=True)
                undergraduateDegreeName=cells[24].get_text(strip=True)
                Percentage_ug=cells[25].get_text(strip=True)
                yearofpassingug=cells[26].get_text(strip=True)

                postgraduateInstituteName=cells[27].get_text(strip=True)
                postgraduateUniversity=cells[28].get_text(strip=True)
                postgraduateDegreeName=cells[29].get_text(strip=True)
                Percentage_pg=cells[30].get_text(strip=True)
                yearofpassingpg=cells[31].get_text(strip=True)

                work_experience=cells[32].get_text(strip=True)
                organization_name=cells[33].get_text(strip=True)
                designation=cells[34].get_text(strip=True)
                time_period_from=cells[35].get_text(strip=True)
                time_period_to=cells[36].get_text(strip=True)

                siblings=cells[37].get_text(strip=True)
                sibling_type=cells[38].get_text(strip=True)
                sibling_age=cells[39].get_text(strip=True)
                sibling_name=cells[40].get_text(strip=True)
                sibling_course=cells[41].get_text(strip=True)
                sibling_institute=cells[42].get_text(strip=True)
                sibling_location=cells[43].get_text(strip=True)

                registration_type=cells[44].get_text(strip=True)

                reference_type=cells[45].get_text(strip=True)
                current_employee_id=cells[46].get_text(strip=True)
                current_employee_name=cells[47].get_text(strip=True)
                relation_with_reference=cells[48].get_text(strip=True)

                ex_employee_name=cells[49].get_text(strip=True)
                ex_employee_id=cells[50].get_text(strip=True)

                walkin_type=cells[51].get_text(strip=True)
                selected=cells[52].get_text(strip=True)

                rows.append([time_stamp,registration_number, name, phone, email, date, fatherName, fatherNumber, fatheroccupation, motherName, Residencial_address, Permanent_address, schoolInstituteName , SchoolUniversity, SchoolDegreeName, Percentage_s, yearofpassings, IntermediateInstituteName, IntermediateUniversity, IntermediateDegreeName , Percentage_i, yearofpassingi, undergraduateInstituteName , undergraduateUniversity , undergraduateDegreeName , Percentage_ug , yearofpassingug , postgraduateInstituteName , postgraduateUniversity , postgraduateDegreeName , Percentage_pg , yearofpassingpg , work_experience , organization_name , designation , time_period_from , time_period_to , siblings , sibling_type, sibling_age, sibling_name , sibling_course , sibling_institute , sibling_location,registration_type, reference_type,current_employee_id, current_employee_name, relation_with_reference,ex_employee_name,ex_employee_id , walkin_type, selected])
        
        # Store the information in a pandas dataframe
        df = pd.DataFrame(rows, columns=['Time Stamp','Registration Number', 'Name', 'Phone', 'Email', 'Date', 'Father Name', 'Father Number', 'Father Occupation', 'Mother Name', 'Residencial Address', 'Permanent Address', 'School Institute Name', 'Stream', 'School Board', 'School Percentage','Year of Passing (School)', 'Intermediate Institute Name', 'Stream' , 'Intermediate Board', 'Intermediate Percentage', 'Year of Passing(Intermediate)', 'Undergraduate Institute Name', 'Undergraduate University' , 'Undergraduate Degree Name' , 'Undergraduate Percentage' , 'Year of Passing(Undergraduate)', 'Post graduate Institute Name', 'Post graduate University' , 'Post graduate Degree Name' , 'Post graduate Percentage' , 'Year of Passing(postgraduate)', 'Prior Work Experience', 'Organization Name', 'Designation','Duration From', 'Duration To', 'Siblings', 'Brother/Sister','Elder/Younger','Name','Study/Working','Oraganization Name','Location', 'Registration Type' , 'Reference Type', 'Employee Id','Employee Name','Relation with Empployee','Ex Employee Name','Ex Employee Id','Walkin Type','Selection Status'])
        
        # Create a BytesIO object to write the CSV file
        output = BytesIO()
        
        # Use the DataFrame to create a CSV file
        df.to_csv(output, index=False)
        
        # Seek to the beginning of the stream
        output.seek(0)
        
        # Send the CSV file as a download
        return send_file(output, as_attachment=True, download_name='users.csv', mimetype='text/csv')
    
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/filter_results', methods=['GET'])
def filter_results():
    session.clear()
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')
    filter_option = request.args.get('filter')
    selection = request.args.get('selection')
    query = {"director_selection": selection}
    today = datetime.now().date()

    if from_date == '':
        from_date = None
    if to_date == '':
        to_date = None

    if from_date is not None and to_date is not None:
        filter_option = None

    if filter_option:
        if filter_option == 'today':
            start_date = today
            end_date = today + timedelta(days=1)
        elif filter_option == 'yesterday':
            start_date = today - timedelta(days=1)
            end_date = today
        elif filter_option == 'week':
            start_date = today - timedelta(days=today.weekday())
            end_date = start_date + timedelta(days=7)
        elif filter_option == 'month':
            start_date = today.replace(day=1)
            next_month = start_date.replace(day=28) + timedelta(days=4)
            end_date = next_month - timedelta(days=next_month.day)
    elif from_date and to_date:
        start_date = datetime.strptime(from_date, '%Y-%m-%d').date()
        end_date = datetime.strptime(to_date, '%Y-%m-%d').date() + timedelta(days=1)
    else:
        start_date = None
        end_date = None

    if start_date and end_date:
        query['director_date_only_str'] = {'$gte': start_date.isoformat(), '$lt': end_date.isoformat()}

    users = list(users_collection.find(query))
    users_list_final = list(users)

    return render_template('director_selection_results.html', users=users_list_final, selection=selection)
@app.route('/filter_results1', methods=['GET'])
def filter_results1():
    session.clear()
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')
    filter_option = request.args.get('filter')
    query = {"register_update": 'registered'}
    today = datetime.now().date()

    if from_date == '':
        from_date = None
    if to_date == '':
        to_date = None

    if from_date is not None and to_date is not None:
        filter_option = None

    if filter_option:
        if filter_option == 'today':
            start_date = today
            end_date = today + timedelta(days=1)
        elif filter_option == 'yesterday':
            start_date = today - timedelta(days=1)
            end_date = today
        elif filter_option == 'week':
            start_date = today - timedelta(days=today.weekday())
            end_date = start_date + timedelta(days=7)
        elif filter_option == 'month':
            start_date = today.replace(day=1)
            next_month = start_date.replace(day=28) + timedelta(days=4)
            end_date = next_month - timedelta(days=next_month.day)
    elif from_date and to_date:
        start_date = datetime.strptime(from_date, '%Y-%m-%d').date()
        end_date = datetime.strptime(to_date, '%Y-%m-%d').date() + timedelta(days=1)
    else:
        start_date = None
        end_date = None

    if start_date and end_date:
        query['registration_date_only_str'] = {'$gte': start_date.isoformat(), '$lt': end_date.isoformat()}

    users = list(users_collection.find(query))
    users_list_final = list(users)

    return render_template('registered_details.html', users=users_list_final)
@app.route('/pdf/<pdf_id>')
def get_pdf(pdf_id):
    # Retrieve the file from GridFS
    grid_out = fs.get(ObjectId(pdf_id))
    pdf_data = grid_out.read()

    # Use send_file to serve the PDF
    return send_file(io.BytesIO(pdf_data), attachment_filename='pdffile.pdf', mimetype='application/pdf')
@app.route('/generate_pdf/<registration_number>', methods=['POST'])
def generate_pdf(registration_number):
    # Retrieve data from MongoDB using the URL parameter
    data = users_collection.find_one({'registration_number': registration_number})

    # data = collection.find_one({'registration_number':registration_number})

    # Check if data is found
    if not data:
        return "No data found for this registration number", 404

    # Generate PDF
    pdf_file = 'generated_pdf.pdf'
    c = canvas.Canvas(pdf_file, pagesize=letter)

    # Example: Inserting data into specific places in PDF
    c.drawString(100, 750, f"Name: {data.get('pname', 'N/A')}")
    c.drawString(100, 730, f"Age: {data.get('age', 'N/A')}")
    # Add more fields as needed

    c.save()

    # Send generated PDF as a downloadable file
    return send_file(pdf_file, as_attachment=True)


if __name__ == '__main__':
    app.run(debug=True)

# Socket server setup